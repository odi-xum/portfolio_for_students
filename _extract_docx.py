import docx
doc = docx.Document('Primer_diplom.docx')

with open('diploma_full.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        style = p.style.name
        if text:
            f.write(f'[{i}] [{style}] {text}\n')
    
    f.write("\n=== TABLES ===\n")
    for ti, table in enumerate(doc.tables, 1):
        f.write(f'\n--- Table {ti} ---\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            f.write(f'  Row {ri}: {" | ".join(cells)}\n')

print("Done")
