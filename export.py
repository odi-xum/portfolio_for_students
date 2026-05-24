"""
Генерация Excel-отчётов с графиками (openpyxl).
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter

from models import Event

# ---------------------------------------------------------------------------
# Стили
# ---------------------------------------------------------------------------
HEADER_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
HEADER_FILL = PatternFill(start_color='3C3836', end_color='3C3836', fill_type='solid')
DATA_FONT = Font(name='Calibri', size=10)
BOLD_FONT = Font(name='Calibri', bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style='thin', color='D5C4A1'),
    right=Side(style='thin', color='D5C4A1'),
    top=Side(style='thin', color='D5C4A1'),
    bottom=Side(style='thin', color='D5C4A1'),
)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
GOLD_FILL = PatternFill(start_color='FBF1C7', end_color='FBF1C7', fill_type='solid')


def _style_header(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _auto_width(ws, min_w=8, max_w=40):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        best = min_w
        for cell in col_cells:
            val = str(cell.value or '')
            length = sum(1.4 if ord(c) > 127 else 1 for c in val)
            if length > best:
                best = length
        ws.column_dimensions[letter].width = min(best + 2, max_w)


def _parse_dt(d, end_of_day=False):
    if not d:
        return None
    fmt = '%Y-%m-%d %H:%M:%S' if end_of_day else '%Y-%m-%d'
    return datetime.strptime(d + (' 23:59:59' if end_of_day else ''), fmt)


def _fio(u):
    return ' '.join(filter(None, [u.last_name, u.first_name, u.patronymic])) or u.username


# ============================================================================
# Sheet 1: Сводка по студентам
# ============================================================================
def _sheet_summary(ws, students, status_filter, date_from, date_to):
    ws.title = 'Сводка'
    headers = ['№', 'ФИО', 'Логин', 'Группа', 'Всего', 'Одобрено',
               'На проверке', 'В комиссии', 'Отклонено', 'Ср. балл']
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(headers))

    row = 2
    for idx, s in enumerate(students, 1):
        q = Event.query.filter(Event.student_id == s.id)
        if status_filter and status_filter != 'all':
            q = q.filter(Event.status == status_filter)
        if date_from:
            q = q.filter(Event.created_at >= _parse_dt(date_from))
        if date_to:
            q = q.filter(Event.created_at <= _parse_dt(date_to, True))

        total = q.count()
        approved = q.filter(Event.status == 'approved').count()
        pending = q.filter(Event.status == 'pending').count()
        disputed = q.filter(Event.status == 'disputed').count()
        rejected = q.filter(Event.status == 'rejected').count()
        scored = q.filter(Event.status == 'approved', Event.score.isnot(None)).all()
        avg = round(sum(e.score for e in scored) / len(scored), 2) if scored else 0

        vals = [idx, _fio(s), s.username, s.group_name or '—',
                total, approved, pending, disputed, rejected, avg]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = DATA_FONT
            cell.alignment = CENTER
            cell.border = THIN_BORDER
            if c == 2:
                cell.alignment = LEFT
        if avg >= 4.5:
            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).fill = GOLD_FILL
        ws.cell(row=row, column=10).number_format = '0.00'
        row += 1

    if students:
        ws.cell(row=row, column=2, value='ИТОГО').font = BOLD_FONT
        for c in [5, 6, 7, 8, 9]:
            cell = ws.cell(row=row, column=c,
                           value=f'=SUM({get_column_letter(c)}2:{get_column_letter(c)}{row - 1})')
            cell.font = BOLD_FONT
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = THIN_BORDER
    _auto_width(ws)


# ============================================================================
# Sheet 2: Детализация мероприятий
# ============================================================================
def _sheet_details(ws, students, status_filter, date_from, date_to):
    ws.title = 'Мероприятия'
    headers = ['№', 'Студент', 'Группа', 'Название', 'Статус', 'Оценка',
               'Комментарий', 'Дата']
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(headers))

    ids = [s.id for s in students]
    q = Event.query.filter(Event.student_id.in_(ids))
    if status_filter and status_filter != 'all':
        q = q.filter(Event.status == status_filter)
    if date_from:
        q = q.filter(Event.created_at >= _parse_dt(date_from))
    if date_to:
        q = q.filter(Event.created_at <= _parse_dt(date_to, True))
    q = q.order_by(Event.student_id, Event.created_at.desc())

    status_ru = {'pending': 'На проверке', 'approved': 'Одобрено',
                 'disputed': 'В комиссии', 'rejected': 'Отклонено'}

    row = 2
    for i, e in enumerate(q.all(), 1):
        stu = e.student
        vals = [i, _fio(stu) if stu else '—',
                stu.group_name if stu else '—',
                e.title, status_ru.get(e.status, e.status),
                e.score if e.score else '—',
                e.curator_comment or '',
                e.created_at.strftime('%d.%m.%Y') if e.created_at else '']
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = DATA_FONT
            cell.alignment = CENTER
            cell.border = THIN_BORDER
            if c in (2, 4, 7):
                cell.alignment = LEFT
        row += 1
    _auto_width(ws)


# ============================================================================
# Sheet 3: Графики
# ============================================================================
def _sheet_charts(ws, students, date_from, date_to):
    ws.title = 'Графики'
    ids = [s.id for s in students]

    def _base():
        q = Event.query.filter(Event.student_id.in_(ids))
        if date_from:
            q = q.filter(Event.created_at >= _parse_dt(date_from))
        if date_to:
            q = q.filter(Event.created_at <= _parse_dt(date_to, True))
        return q

    # ---- Таблица статусов ----
    ws.cell(row=1, column=1, value='Статус').font = HEADER_FONT
    ws.cell(row=1, column=2, value='Кол-во').font = HEADER_FONT
    _style_header(ws, 1, 2)

    labels = ['Одобрено', 'На проверке', 'В комиссии', 'Отклонено']
    keys = ['approved', 'pending', 'disputed', 'rejected']
    colors = ['A9B665', 'D8A657', 'D3869B', 'EA6962']
    bq = _base()
    for i, (lab, key) in enumerate(zip(labels, keys), 2):
        cnt = bq.filter(Event.status == key).count()
        ws.cell(row=i, column=1, value=lab).font = DATA_FONT
        ws.cell(row=i, column=2, value=cnt).font = DATA_FONT
        for c in (1, 2):
            ws.cell(row=i, column=c).border = THIN_BORDER
            ws.cell(row=i, column=c).alignment = CENTER
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 12

    total = bq.count()
    pie = PieChart()
    pie.title = f'Статусы мероприятий (всего: {total})'
    pie.width = 18
    pie.height = 14
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    for i, c in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = c
        pie.series[0].data_points.append(pt)
    ws.add_chart(pie, 'D1')

    # ---- Таблица среднего балла ----
    ws.cell(row=1, column=7, value='Студент').font = HEADER_FONT
    ws.cell(row=1, column=8, value='Ср. балл').font = HEADER_FONT
    _style_header(ws, 1, 8)

    scores = []
    for s in students:
        q = Event.query.filter(
            Event.student_id == s.id, Event.status == 'approved',
            Event.score.isnot(None)
        )
        if date_from:
            q = q.filter(Event.created_at >= _parse_dt(date_from))
        if date_to:
            q = q.filter(Event.created_at <= _parse_dt(date_to, True))
        vals = [e.score for e in q.all()]
        avg = round(sum(vals) / len(vals), 2) if vals else 0
        scores.append((_fio(s), avg))

    scores.sort(key=lambda x: x[1], reverse=True)
    bar_row = 2
    for i, (fio, avg) in enumerate(scores[:15], 2):
        ws.cell(row=i, column=7, value=fio).font = DATA_FONT
        ws.cell(row=i, column=8, value=avg).font = DATA_FONT
        ws.cell(row=i, column=8).number_format = '0.00'
        for c in (7, 8):
            ws.cell(row=i, column=c).border = THIN_BORDER
        bar_row = i

    ws.column_dimensions['G'].width = 22
    ws.column_dimensions['H'].width = 12

    bar = BarChart()
    bar.type = 'col'
    bar.title = 'Средний балл (топ-15)'
    bar.y_axis.title = 'Средний балл'
    bar.width = 24
    bar.height = 14
    bar.add_data(Reference(ws, min_col=8, min_row=1, max_row=bar_row), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=7, min_row=2, max_row=bar_row))
    bar.series[0].graphicalProperties.solidFill = '7DAEA3'
    ws.add_chart(bar, 'D18')


# ============================================================================
# ПУБЛИЧНАЯ ФУНКЦИЯ
# ============================================================================
def generate_report(students, include_charts=True, include_details=True,
                    status_filter='all', date_from=None, date_to=None):
    """
    Генерирует Excel-файл с отчётом. Возвращает BytesIO.

    Параметры:
        students — список User
        include_charts — лист с графиками
        include_details — лист с детализацией
        status_filter — 'all' | 'pending' | 'approved' | 'disputed' | 'rejected'
        date_from — 'YYYY-MM-DD' или None
        date_to   — 'YYYY-MM-DD' или None
    """
    wb = Workbook()
    _sheet_summary(wb.active, students, status_filter, date_from, date_to)

    if include_details:
        _sheet_details(wb.create_sheet(), students, status_filter, date_from, date_to)

    if include_charts:
        _sheet_charts(wb.create_sheet(), students, date_from, date_to)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
