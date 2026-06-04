"""
Импорт студентов из Excel-файла.
"""
import io
from openpyxl import load_workbook
from database import SessionLocal
from models import User, Group
from dependencies import hash_password


def import_students_from_excel(file_storage):
    """
    Читает .xlsx, создаёт студентов. Возвращает (imported, skipped, errors).
    """
    imported = 0; skipped = 0; errors = []
    wb = load_workbook(file_storage, read_only=True)
    ws = wb.active
    if ws is None: return 0, 0, ['Файл не содержит листов']

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return 0, 0, ['Файл пуст']

    col_map = {}
    for i, h in enumerate(header):
        if h is None: continue
        hl = str(h).strip().lower()
        if 'фамили' in hl: col_map['last'] = i
        elif 'имя' in hl and 'отчеств' not in hl: col_map['first'] = i
        elif 'отчеств' in hl: col_map['patr'] = i
        elif 'групп' in hl: col_map['group'] = i
        elif 'логин' in hl: col_map['login'] = i
        elif 'парол' in hl: col_map['password'] = i

    if 'login' not in col_map:
        return 0, 0, ['Не найдена колонка "Логин".']

    db = SessionLocal()
    try:
        for row_idx, row in enumerate(rows_iter, 2):
            try:
                login = _val(row, col_map.get('login'))
                if not login: skipped += 1; continue
                if db.query(User).filter(User.username == str(login)).first():
                    skipped += 1; continue
                group_name = _val(row, col_map.get('group'))
                group_id = None
                if group_name:
                    grp = db.query(Group).filter(Group.name == group_name).first()
                    if grp: group_id = grp.id
                db.add(User(
                    username=str(login),
                    password_hash=hash_password(_val(row, col_map.get('password'), '111')),
                    role='student',
                    last_name=_val(row, col_map.get('last')),
                    first_name=_val(row, col_map.get('first')),
                    patronymic=_val(row, col_map.get('patr')),
                    group_name=group_name, group_id=group_id,
                ))
                imported += 1
            except Exception as e:
                errors.append(f'Строка {row_idx}: {e}'); skipped += 1
        db.commit()
    finally:
        db.close()
    return imported, skipped, errors


def _val(row, idx, default=None):
    if idx is None or idx >= len(row): return default
    v = row[idx]
    if v is None: return default
    return str(v).strip()
